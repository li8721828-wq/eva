from docx import Document
from docx.shared import Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.section import WD_SECTION
from pathlib import Path

OUT = Path("docs")

def base(title, subtitle):
    d=Document(); sec=d.sections[0]; sec.top_margin=Pt(54); sec.bottom_margin=Pt(54); sec.left_margin=Pt(68); sec.right_margin=Pt(68)
    styles=d.styles
    styles["Normal"].font.name="Microsoft YaHei"; styles["Normal"].font.size=Pt(10.5); styles["Normal"].font.color.rgb=RGBColor(45,55,72)
    for n,size,color in [("Title",25,(25,45,75)),("Heading 1",16,(25,90,120)),("Heading 2",12,(25,90,120))]:
        styles[n].font.name="Microsoft YaHei"; styles[n].font.size=Pt(size); styles[n].font.bold=True; styles[n].font.color.rgb=RGBColor(*color)
    p=d.add_paragraph(style="Title"); p.add_run(title)
    p=d.add_paragraph(subtitle); p.runs[0].font.color.rgb=RGBColor(100,110,125)
    d.add_paragraph("版本 0.1.206  |  更新日期 2026 年 9 月 15 日")
    return d
def bullets(d, items):
    for x in items: d.add_paragraph(x, style="List Bullet")

d=base("Eva 今日功能更新说明","面向使用者的更新概览")
d.add_heading("本次更新解决了什么",1)
d.add_paragraph("本次更新围绕 Agent 执行可靠性、工具使用安全性、本地服务能力和聊天体验展开。Eva 现在能够更清晰地记录任务执行过程，在调用工具前保留审批入口，并为后续外部集成和受控执行提供基础。")
d.add_heading("主要变化",1)
bullets(d,["新增本地 App Server 基础能力，可为外部客户端提供任务、状态和事件访问入口。","新增工具审批策略与审批卡片，工具执行可以进入明确的允许、拒绝或待处理状态。","接入 Rust 沙箱层，为终端和文件类操作提供更清晰的执行边界。","补充 Agent 运行事件日志，记录运行、轮次、模型调用和工具调用的开始与完成状态。","增强中断与恢复信息，异常中断后可以识别未完成的运行过程。","优化聊天消息列表、滚动恢复、欢迎页和设置界面，减少输出混乱并改善操作反馈。"])
d.add_heading("验证结果",1)
d.add_paragraph("类型检查已通过。与本次功能相关的核心单元测试共 35 项，全部通过。Electron 依赖已恢复，当前项目可以继续进行开发和打包验证。")
d.save(OUT/"2026-09-15-Eva-功能更新说明.docx")

d=base("Eva 今日技术变更记录","面向开发与维护的实现记录")
d.add_heading("变更范围",1)
d.add_paragraph("本次变更集中在运行时生命周期、服务边界、审批控制、沙箱执行和渲染层体验。版本号保持为 0.1.206。")
d.add_heading("运行时与事件链路",1)
bullets(d,["AgentRunner 写入 run_started、turn_started、model_call、tool_call、turn_completed 和 run_completed 事件。","增加运行事件存储和恢复摘要，用于定位最后完成轮次、未完成轮次与中断状态。","补充中断上下文处理，使调度层能够获得运行恢复所需的信息。"])
d.add_heading("服务与安全边界",1)
bullets(d,["新增 App Server IPC 与本地服务目录，为未来 HTTP 或外部客户端接入预留统一入口。","新增工具审批策略，集中判断工具是否需要用户确认。","新增 Rust sandbox backend、profile builder 与验证脚本，建立受控执行的基础结构。"])
d.add_heading("前端与 IPC",1)
bullets(d,["扩展共享 IPC channel、contract 和任务类型，保持主进程、预加载层和渲染层的类型一致。","新增 ToastViewport、ToolApprovalCard、AppServerPanel 与滚动位置恢复 Hook。","整理消息列表和对话状态管理，改善流式输出、执行反馈和界面布局。"])
d.add_heading("验证与后续工作",1)
d.add_paragraph("npm run typecheck 已通过；核心测试 35 项全部通过。后续应继续补齐 model_call 与 tool_call 的更多异常路径测试，并完成完整打包验证和远程提交。")
d.save(OUT/"2026-09-15-Eva-技术变更记录.docx")
